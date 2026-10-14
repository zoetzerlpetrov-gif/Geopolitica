import json, re, sys, io
sys.path.insert(0, "ingest")
import fuentes as F
import openpyxl
B = "https://www.gob.mx"
for f in ["/cms/uploads/attachment/file/1105801/Valor_de_inmuebles_valuados_para_tipo_de_vivienda_casa-habitaci_n.xlsx",
          "/cms/uploads/attachment/file/1105803/Valor_de_inmuebles_valuados_para_tipo_de_vivienda_departamento.xlsx"]:
    b = F.get(B + f, timeout=90); print("\n#####", f, len(b))
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    for ws in wb.worksheets:
        print("== hoja", ws.title, ws.max_row, ws.max_column)
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i < 14 or i % 40 == 0: print(i, [c for c in row[:40]])
            if i > 400: break
b = F.get(B + "/shf/documentos/indicadores-sociedad-hipotecaria-federal", timeout=60).decode("utf-8", "replace")
txt = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", b)); i = txt.find("Clase Valor"); print(txt[i-200:i+2500])
b = F.get(B + "/shf/es/archivo/documentos", timeout=60).decode("utf-8", "replace")
print([h for h in re.findall(r'href="([^"]+)"', b) if "indice" in h.lower() or "precios" in h.lower()])
for u in ["/shf/documentos/indice-shf-de-precios-de-la-vivienda-en-mexico-primer-trimestre-de-2026"]:
    try:
        b = F.get(B + u, timeout=60).decode("utf-8", "replace"); print("PAGE", u, [h for h in re.findall(r'href="([^"]+)"', b) if "attachment" in h])
    except Exception as e: print("PAGE", u, e)
for t in ["List_of_paranormal_events_reported_in_Wiltshire,_England", "List_of_space_debris_fall_incidents", "Hoia-Baciu_Forest", "Snippy_(horse)", "Athena_rocket", "Ruta_del_OVNI"]:
    for lang in ["en", "es"]:
        try:
            d = json.loads(F.get(f"https://{lang}.wikipedia.org/api/rest_v1/page/summary/{t}", timeout=30))
            print("WIKI", lang, t, d.get("coordinates"), d["content_urls"]["desktop"]["page"]); break
        except Exception as e: print("WIKI", lang, t, e)
